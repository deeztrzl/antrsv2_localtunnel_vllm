#!/usr/bin/env python3
"""
claim_verifier.py  (versi vLLM)
===============================
Agent verifikator klaim RS & JKK, jalan lewat command line.

Alur kerja (3 tahap, makanya disebut "agent" bukan sekadar 1x panggilan LLM):

    1. EKSTRAKSI  : baca semua dokumen klaim (invoice, resume medis, PLKK, dst)
                     -> minta LLM mengekstrak daftar item tagihan jadi JSON
    2. RETRIEVAL  : untuk tiap item tagihan, cari tarif pembanding paling mirip
                     di database `dokumen_chunk` (vector search, RS yang sama)
    3. ANALISIS   : kirim dokumen klaim + item tagihan + tarif pembanding hasil
                     retrieval ke LLM, memakai system prompt verifikator,
                     minta hasil akhir dalam JSON terstruktur

    -> disimpan sebagai laporan_<kode_rs>_<tanggal>.md, .xlsx, dan .csv

LLM dipanggil lewat server vLLM (OpenAI-compatible API), misalnya yang
diakses lewat tunnel.

PENTING (batas tanggung jawab tool ini):
    Ini adalah decision-support tool. Keputusan akhir klaim tetap harus
    diperiksa dan ditandatangani oleh verifikator manusia. Skrip ini tidak
    boleh dipakai untuk membuat keputusan otomatis tanpa review manusia.

Cara pakai:
    Salin .env.example jadi .env lalu isi LLM_BASE_URL, LLM_MODEL, (opsional) LLM_API_KEY.
    URL tunnel berubah? Cukup edit .env, tidak perlu ubah kode.

    python claim_verifier.py \
        --kode-rs RS001 \
        --claim-docs invoice.pdf resume_medis.pdf plkk.pdf \
        --dsn "postgresql://admin:pass@localhost:5432/poc_rdbms"

Lihat requirements.txt untuk dependensi (ganti google-genai dengan openai).
File ini butuh pdf_to_db.py (fungsi extract_pdf & build_markdown_and_json)
ada di folder yang sama.
"""

import argparse
import base64
import collections
import csv
import json
import os
import re
import sys
import threading
import time
from datetime import date, datetime
from pathlib import Path

import fitz  # PyMuPDF, dipakai untuk render halaman jadi gambar
import numpy as np
import psycopg2
from dotenv import load_dotenv
from openai import BadRequestError, OpenAI
from openpyxl import Workbook
from openpyxl.styles import Font
from sentence_transformers import SentenceTransformer

# Reuse fungsi dari pipeline ingest supaya cara baca PDF konsisten
from pdf_to_db import EMBEDDING_MODEL_NAME, build_markdown_and_json, extract_pdf

# ---------------------------------------------------------------------------
# Konfigurasi
# ---------------------------------------------------------------------------

# Muat variabel dari file .env (di folder skrip ini, atau folder tempat skrip dijalankan).
# Harus dipanggil SEBELUM os.environ.get() di bawah. Nilai di .env MENANG atas env var shell
# (override=True), supaya tidak ada nilai lama yang tertinggal di shell diam-diam menimpa .env.
load_dotenv(Path(__file__).parent / ".env", override=True)

# vLLM biasanya hanya men-serve satu model, jadi kedua tahap memakai model yang sama.
# Isi lewat env var LLM_MODEL (lihat nama persisnya di: curl <base_url>/models).
DEFAULT_MODEL = "Qwen/Qwen2.5-VL-7B-Instruct"  # sesuaikan dengan model yang di-serve vLLM
# Opsional: model berbeda per tahap (mis. Gemini flash-lite untuk ekstraksi, flash untuk analisis).
MODEL_EKSTRAKSI = os.environ.get("LLM_MODEL_EKSTRAKSI") or os.environ.get("LLM_MODEL", DEFAULT_MODEL)
MODEL_ANALISIS = os.environ.get("LLM_MODEL_ANALISIS") or os.environ.get("LLM_MODEL", DEFAULT_MODEL)

# Batas token output. Harus lebih kecil dari (max-model-len server - token prompt).
# GPU Colab gratis (T4) biasanya context-nya kecil, jadi bisa diatur lewat .env tanpa edit kode.
MAX_OUTPUT_EKSTRAKSI = int(os.environ.get("LLM_MAX_OUTPUT_EKSTRAKSI", 4096))
MAX_OUTPUT_ANALISIS = int(os.environ.get("LLM_MAX_OUTPUT_ANALISIS", 8192))

# Ukuran potongan teks klaim per panggilan ekstraksi (karakter). ~4000 karakter = ~1500 token.
EKSTRAKSI_CHUNK_CHARS = int(os.environ.get("LLM_EKSTRAKSI_CHUNK_CHARS", 4000))

# --- Analisis berjenjang (untuk server dengan context window kecil) ---
# Isi LLM_MAX_CONTEXT sama dengan --max-model-len di server vLLM.
LLM_MAX_CONTEXT = int(os.environ.get("LLM_MAX_CONTEXT", 32768))
# auto = pakai analisis penuh kalau muat, selain itu berjenjang. Bisa dipaksa: penuh / berjenjang.
LLM_ANALISIS_MODE = os.environ.get("LLM_ANALISIS_MODE", "auto").lower()
ANALISIS_BATCH_ITEM = int(os.environ.get("LLM_ANALISIS_BATCH_ITEM", 10))   # item per panggilan
TARIF_CHUNK_CHARS = int(os.environ.get("LLM_TARIF_CHUNK_CHARS", 800))     # potong tiap tarif pembanding
MAX_OUTPUT_BATCH = int(os.environ.get("LLM_MAX_OUTPUT_BATCH", 2000))
MAX_OUTPUT_FINAL = int(os.environ.get("LLM_MAX_OUTPUT_FINAL", 4000))
CHARS_PER_TOKEN = 2.5  # taksiran konservatif untuk teks Indonesia + JSON
TOKEN_PER_GAMBAR = int(os.environ.get("LLM_TOKEN_PER_GAMBAR", 3000))  # taksiran token 1 gambar halaman (zoom 2.0)

# Jeda minimum antar panggilan LLM (detik). Berguna untuk free tier Gemini yang membatasi request per menit
# (mis. 5 RPM -> isi 13). 0 = tanpa jeda.
JEDA_MIN_DETIK = float(os.environ.get("LLM_MIN_JEDA_DETIK", 0))

# Batas keras ukuran PROMPT per panggilan (token, taksiran). Panggilan yang melebihinya ditolak SEBELUM dikirim,
# jadi prompt yang membengkak tidak menghabiskan kuota (mis. TPM free tier Gemini). 0 = tanpa batas tambahan.
BATAS_PROMPT_TOKEN = int(os.environ.get("LLM_MAX_PROMPT_TOKEN", 0))

TOP_K_TARIF = 3  # jumlah kandidat tarif pembanding diambil per item tagihan

SYSTEM_PROMPT_PATH = Path(__file__).parent / "system_prompt_verifikator.md"


# ---------------------------------------------------------------------------
# Util kecil
# ---------------------------------------------------------------------------

# ---------------------------------------------------------------------------
# Retry helper -- koneksi lewat tunnel kadang putus-nyambung,
# jangan sampai kerja 73 item hilang cuma karena gangguan sesaat.
# ---------------------------------------------------------------------------

_throttle_lock = threading.Lock()
_terakhir_panggil = [0.0]


def _tunggu_jeda() -> None:
    """Jaga jarak minimal JEDA_MIN_DETIK antar panggilan (aman dipakai dari beberapa thread job)."""
    if JEDA_MIN_DETIK <= 0:
        return
    with _throttle_lock:
        tunggu = _terakhir_panggil[0] + JEDA_MIN_DETIK - time.time()
        if tunggu > 0:
            time.sleep(tunggu)
        _terakhir_panggil[0] = time.time()


def panggil_llm(client: OpenAI, max_retries: int = 6, delay_awal: int = 5, label: str = "LLM call", usage_log: list = None, **kwargs):
    """Wrapper client.chat.completions.create dengan retry + exponential backoff.
    Kalau usage_log (list) diberikan, catat durasi & jumlah token panggilan ini ke situ."""
    delay = delay_awal
    mulai = time.time()

    # Pengaman: max_tokens yang terlalu besar (mis. sama dengan seluruh context window) membuat
    # server menolak request walau prompt-nya kecil. Kecilkan otomatis sesuai sisa context.
    if kwargs.get("max_tokens"):
        taksiran_prompt = 0
        for m in kwargs.get("messages", []):
            isi = m.get("content", "")
            if isinstance(isi, str):
                taksiran_prompt += _taksir_token(isi)
            else:
                for bagian in isi:
                    if bagian.get("type") == "text":
                        taksiran_prompt += _taksir_token(bagian.get("text", ""))
                    else:
                        taksiran_prompt += TOKEN_PER_GAMBAR
        if BATAS_PROMPT_TOKEN and taksiran_prompt > BATAS_PROMPT_TOKEN:
            raise RuntimeError(
                f"[{label}] prompt ~{taksiran_prompt} token melebihi LLM_MAX_PROMPT_TOKEN={BATAS_PROMPT_TOKEN}; "
                f"tidak dikirim. Periksa bagian [UKURAN PROMPT] di log untuk melihat komponen yang membengkak."
            )
        batas = LLM_MAX_CONTEXT - taksiran_prompt - 300  # 300 = cadangan untuk template chat
        if batas < 256:
            raise RuntimeError(
                f"Prompt terlalu besar untuk context {LLM_MAX_CONTEXT} token (perkiraan ~{taksiran_prompt} token). "
                f"Kecilkan LLM_EKSTRAKSI_CHUNK_CHARS / LLM_ANALISIS_BATCH_ITEM, atau naikkan --max-model-len server."
            )
        if kwargs["max_tokens"] > batas:
            print(f"      [INFO] [{label}] max_tokens {kwargs['max_tokens']} diturunkan ke {batas} "
                  f"(context {LLM_MAX_CONTEXT}, prompt ~{taksiran_prompt} token).")
            kwargs["max_tokens"] = batas

    for percobaan in range(1, max_retries + 1):
        try:
            _tunggu_jeda()
            resp = client.chat.completions.create(**kwargs)
            durasi = time.time() - mulai

            usage = resp.usage
            prompt_tok = usage.prompt_tokens if usage else None
            output_tok = usage.completion_tokens if usage else None
            total_tok = usage.total_tokens if usage else None

            print(f"      [{label}] {durasi:.1f}s -- token: prompt {prompt_tok}, output {output_tok}, total {total_tok}")
            if resp.choices and resp.choices[0].finish_reason == "length":
                print("      [WARN] Output terpotong (finish_reason=length). Naikkan max_tokens / cek --max-model-len server.")
            if usage_log is not None:
                usage_log.append({
                    "label": label, "durasi_detik": round(durasi, 1),
                    "prompt_token": prompt_tok, "output_token": output_tok, "total_token": total_tok,
                })
            return resp
        except BadRequestError as e:
            # 400 = request-nya sendiri yang bermasalah (mis. melebihi context length). Retry sia-sia.
            raise RuntimeError(f"Server LLM menolak request (400): {e}") from e
        except Exception as e:
            if percobaan == max_retries:
                raise
            # APIConnectionError cuma bilang "Connection error."; penyebab aslinya (DNS, refused,
            # timeout, SSL, connection reset) ada di __cause__.
            sebab = f" | sebab: {e.__cause__!r}" if getattr(e, "__cause__", None) else ""
            print(f"      [WARN] Panggilan LLM gagal ({e.__class__.__name__}: {e}{sebab}), coba lagi dalam {delay}s (percobaan {percobaan}/{max_retries})...")
            time.sleep(delay)
            delay = min(delay * 2, 60)


def buat_llm_client(base_url: str, api_key: str = "EMPTY") -> OpenAI:
    """Client OpenAI-compatible untuk vLLM lewat tunnel. Dipakai oleh CLI maupun api.py."""
    return OpenAI(
        base_url=base_url,
        api_key=api_key or "EMPTY",
        timeout=900,  # generate panjang lewat tunnel butuh timeout longgar
        max_retries=0,  # retry sudah ditangani panggil_llm
        # localtunnel (loca.lt) menampilkan halaman "reminder" untuk request tanpa header ini,
        # sehingga API call malah dapat HTML. ngrok free pakai header yang kedua.
        # Keduanya tidak berbahaya kalau tunnel-nya jenis lain.
        default_headers={
            "bypass-tunnel-reminder": "true",
            "ngrok-skip-browser-warning": "1",
        },
    )


def muat_ulang_konfigurasi() -> None:
    """
    Baca ulang file .env (nilai di .env MENANG atas env var shell) dan perbarui semua konfigurasi LLM
    tanpa restart proses. Dipanggil api.py di awal tiap job, jadi URL tunnel baru cukup diedit di .env.
    """
    global MODEL_EKSTRAKSI, MODEL_ANALISIS, MAX_OUTPUT_EKSTRAKSI, MAX_OUTPUT_ANALISIS, EKSTRAKSI_CHUNK_CHARS
    global LLM_MAX_CONTEXT, LLM_ANALISIS_MODE, ANALISIS_BATCH_ITEM, TARIF_CHUNK_CHARS
    global MAX_OUTPUT_BATCH, MAX_OUTPUT_FINAL, TOKEN_PER_GAMBAR, BATAS_PROMPT_TOKEN
    env_path = Path(__file__).parent / ".env"
    load_dotenv(env_path if env_path.exists() else None, override=True)
    e = os.environ.get
    MODEL_EKSTRAKSI = e("LLM_MODEL_EKSTRAKSI") or e("LLM_MODEL", DEFAULT_MODEL)
    MODEL_ANALISIS = e("LLM_MODEL_ANALISIS") or e("LLM_MODEL", DEFAULT_MODEL)
    MAX_OUTPUT_EKSTRAKSI = int(e("LLM_MAX_OUTPUT_EKSTRAKSI", 4096))
    MAX_OUTPUT_ANALISIS = int(e("LLM_MAX_OUTPUT_ANALISIS", 8192))
    EKSTRAKSI_CHUNK_CHARS = int(e("LLM_EKSTRAKSI_CHUNK_CHARS", 4000))
    LLM_MAX_CONTEXT = int(e("LLM_MAX_CONTEXT", 32768))
    LLM_ANALISIS_MODE = e("LLM_ANALISIS_MODE", "auto").lower()
    ANALISIS_BATCH_ITEM = int(e("LLM_ANALISIS_BATCH_ITEM", 10))
    TARIF_CHUNK_CHARS = int(e("LLM_TARIF_CHUNK_CHARS", 800))
    MAX_OUTPUT_BATCH = int(e("LLM_MAX_OUTPUT_BATCH", 2000))
    MAX_OUTPUT_FINAL = int(e("LLM_MAX_OUTPUT_FINAL", 4000))
    TOKEN_PER_GAMBAR = int(e("LLM_TOKEN_PER_GAMBAR", 3000))
    BATAS_PROMPT_TOKEN = int(e("LLM_MAX_PROMPT_TOKEN", 0))


def cek_server_llm(client: OpenAI) -> list:
    """
    Cek cepat (maks ~20 detik) bahwa server LLM terjangkau SEBELUM pekerjaan berat (baca PDF, ekstraksi) dimulai.
    Tanpa ini, tunnel yang mati baru ketahuan setelah beberapa menit dan 6x retry.
    Return daftar id model di server. Raise RuntimeError dengan pesan yang bisa ditindaklanjuti.
    """
    url = str(getattr(client, "base_url", ""))
    try:
        daftar = [m.id for m in client.with_options(timeout=20).models.list().data]
    except Exception as e:
        if e.__class__.__name__ in ("AuthenticationError", "PermissionDeniedError"):
            raise RuntimeError(f"Server LLM di {url} menolak API key ({e.__class__.__name__}). Cek LLM_API_KEY di .env.") from e
        sebab = f" | sebab: {e.__cause__!r}" if getattr(e, "__cause__", None) else ""
        raise RuntimeError(
            f"Server LLM tidak terjangkau di {url} ({e.__class__.__name__}{sebab}). Kalau memakai tunnel gratis "
            f"(trycloudflare/ngrok/loca.lt), URL-nya kemungkinan sudah berubah atau mati: minta URL baru, "
            f"ubah LLM_BASE_URL di .env, lalu coba lagi."
        ) from e
    for m in {MODEL_EKSTRAKSI, MODEL_ANALISIS}:
        if daftar and m not in daftar and not any(x.endswith("/" + m) for x in daftar):
            print(f"      [WARN] Model '{m}' tidak ada di daftar server ({', '.join(daftar[:5])}). Cek LLM_MODEL di .env.")
    return daftar


def teks_dari(resp) -> str:
    """Ambil teks jawaban dari response chat completion."""
    if not resp.choices:
        return ""
    return resp.choices[0].message.content or ""


def parse_json_response(text: str):
    """
    Model open-source (terutama yang punya mode thinking seperti Qwen3) sering menyisipkan
    <think>...</think>, teks pembuka, atau pembungkus ```json. Bersihkan dulu, lalu ambil
    dari kurung pertama sampai kurung terakhir.
    """
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.DOTALL).strip()
    if text.startswith("```"):
        text = text.strip("`")
        text = text.split("\n", 1)[1] if "\n" in text else text
        if text.lower().startswith("json"):
            text = text[4:]
    kandidat = [i for i in (text.find("{"), text.find("[")) if i != -1]
    mulai = min(kandidat) if kandidat else 0
    akhir = max(text.rfind("}"), text.rfind("]"))
    if akhir == -1:
        akhir = len(text) - 1
    return json.loads(text[mulai:akhir + 1])


def _salvage_json_items(text: str) -> list:
    """
    Jika JSON output dari LLM terpotong di tengah jalan (finish_reason=length),
    selamatkan objek-objek item yang sudah terbentuk lengkap sebelum titik potong.
    """
    clean = re.sub(r"<think>.*?</think>", "", text, flags=re.DOTALL).strip()
    if clean.startswith("```"):
        clean = clean.strip("`")
        clean = clean.split("\n", 1)[1] if "\n" in clean else clean
        if clean.lower().startswith("json"):
            clean = clean[4:].strip()

    # Strategi 1: Coba tutup kurung array secara paksa di objek valid terakhir
    pos_items = clean.find('"items"')
    if pos_items != -1:
        kurung_buka = clean.find('[', pos_items)
        if kurung_buka != -1:
            terakhir_tutup = clean.rfind('}')
            if terakhir_tutup > kurung_buka:
                potongan = clean[:terakhir_tutup + 1] + "]}"
                try:
                    data = json.loads(potongan)
                    if isinstance(data, dict) and isinstance(data.get("items"), list) and data["items"]:
                        return data["items"]
                except Exception:
                    pass

    # Strategi 2: Regex extraction tiap blok {...} yang memiliki key "item"
    salvaged = []
    pola = re.compile(r'\{[^{}]*?"item"\s*:[^{}]*?\}', re.DOTALL)
    for m in pola.finditer(clean):
        try:
            obj = json.loads(m.group(0))
            if isinstance(obj, dict) and obj.get("item"):
                salvaged.append(obj)
        except Exception:
            continue
    return salvaged


def halaman_kemungkinan_gambar(page: dict) -> bool:
    """
    Deteksi halaman yang perlu dibaca ulang lewat vision (dikirim sebagai gambar ke LLM):
    1. Benar-benar kosong (scan murni).
    2. Teks sangat sedikit (< 500 karakter) dan tidak memiliki tabel terstruktur pdfplumber/Camelot.
    3. Halaman rincian tagihan/billing medis namun tabel terstruktur gagal dideteksi (tabel tanpa garis).
    """
    text = page.get("text", "").strip()
    ada_isi_tabel = any(
        any(str(c).strip() for c in row) for tbl in page.get("tables", []) for row in tbl.get("rows", [])
    )
    if not text and not ada_isi_tabel:
        return True  # kosong total (scan murni)
    if len(text) < 500 and not ada_isi_tabel:
        return True  # teks minim / kop surat / noise OCR tanpa tabel terstruktur

    # Deteksi tabel billing tanpa garis (borderless table) yang gagal diekstrak Camelot
    text_lower = text.lower()
    pola_billing = ("rincian biaya", "detail tagihan", "kwitansi", "billing", "detail layanan", "rekap biaya")
    if any(k in text_lower for k in pola_billing) and not ada_isi_tabel:
        return True

    return False


def render_halaman_jadi_gambar(pdf_path: str, page_number: int, zoom: float = 2.0) -> bytes:
    """Render satu halaman PDF (1-indexed) jadi PNG bytes, buat dikirim ke LLM (vision)."""
    doc = fitz.open(pdf_path)
    try:
        page = doc[page_number - 1]
        mat = fitz.Matrix(zoom, zoom)
        pix = page.get_pixmap(matrix=mat)
        return pix.tobytes("png")
    finally:
        doc.close()


def _kompres_teks(md: str) -> str:
    """
    Buang pemborosan token tanpa mengubah isi: baris pemisah tabel Markdown (| --- | --- |), spasi berlebih,
    dan baris kosong beruntun. Baris identik BERURUTAN sengaja TIDAK dibuang: itu bisa penagihan ganda yang sah.
    """
    hasil = []
    for baris in md.splitlines():
        b = re.sub(r"[ \t]+", " ", baris).rstrip()
        if "---" in b and set(b.replace(" ", "")) <= set("|-:"):
            continue
        if not b.strip() and (not hasil or not hasil[-1].strip()):
            continue
        hasil.append(b)
    return "\n".join(hasil)


def _saring_teks_halaman(page: dict) -> int:
    """
    extract_text() pdfplumber ikut mengambil isi tabel, padahal tabel yang sama ditulis lagi sebagai tabel Markdown
    (hasil Camelot): isi tabel masuk prompt DUA kali. Baris teks dibuang hanya kalau persis sama (setelah
    dinormalisasi) dengan salah satu baris tabel, dan sebanyak-banyaknya sejumlah kemunculannya di tabel, jadi baris
    ganda yang sah tidak ikut hilang. Return jumlah baris yang dibuang.
    """
    if not page.get("tables"):
        return 0
    baris_tabel = [k for tbl in page["tables"] for row in tbl["rows"]
                   for k in [_norm_nama("".join(str(c) for c in row))] if k]
    if not baris_tabel:
        return 0
    terpakai, baru, buang = collections.Counter(), [], 0
    for baris in page["text"].split("\n"):
        k = _norm_nama(baris)
        if len(k) >= 12 and terpakai[k] < sum(1 for t in baris_tabel if k in t):
            terpakai[k] += 1
            buang += 1
            continue
        baru.append(baris)
    page["text"] = "\n".join(baru)
    return buang


def baca_semua_dokumen_klaim(paths: list) -> tuple:
    """
    Gabungkan hasil ekstraksi (teks + tabel) semua file klaim jadi satu blok teks,
    dan kumpulkan gambar untuk halaman yang kemungkinan hasil scan (tidak ada teks
    yang terbaca), supaya bisa dibaca langsung oleh LLM lewat vision.

    Return: (teks_gabungan, list_gambar_png_bytes)
    """
    bagian = []
    gambar_halaman_gambar = []

    for p in paths:
        p = Path(p)
        if not p.exists():
            sys.exit(f"File tidak ditemukan: {p}")
        print(f"      membaca: {p.name}")
        extracted = extract_pdf(str(p))
        # Deteksi halaman scan DULU (berdasarkan teks asli), baru teks dirapikan.
        halaman_gambar = [
            pg["page_number"] for pg in extracted["pages"] if halaman_kemungkinan_gambar(pg)
        ]
        n_duplikat = sum(_saring_teks_halaman(pg) for pg in extracted["pages"])
        md, _ = build_markdown_and_json(extracted, nama_dokumen=p.stem)
        panjang_awal = len(md)
        md = _kompres_teks(md)
        if n_duplikat or len(md) < panjang_awal:
            print(f"      -> {p.name}: {n_duplikat} baris teks duplikat tabel dibuang; {panjang_awal} -> {len(md)} karakter")
        bagian.append(f"\n\n=== DOKUMEN: {p.name} ===\n{md}")
        if halaman_gambar:
            print(f"      -> {len(halaman_gambar)} halaman di {p.name} terdeteksi sebagai gambar/scan, akan dibaca lewat vision")
            for no in halaman_gambar:
                gambar_halaman_gambar.append(render_halaman_jadi_gambar(str(p), no))

    return "\n".join(bagian), gambar_halaman_gambar


# ---------------------------------------------------------------------------
# Tahap 1: Ekstraksi item tagihan
# ---------------------------------------------------------------------------

def _pecah_teks(teks: str, batas: int) -> list:
    """Pecah teks jadi potongan <= batas karakter, memotong di batas baris."""
    potongan, buf = [], ""
    for baris in teks.splitlines(keepends=True):
        while len(baris) > batas:  # satu baris super panjang
            if buf:
                potongan.append(buf)
                buf = ""
            potongan.append(baris[:batas])
            baris = baris[batas:]
        if buf and len(buf) + len(baris) > batas:
            potongan.append(buf)
            buf = ""
        buf += baris
    if buf.strip():
        potongan.append(buf)
    return potongan


def _ekstrak_satu(client: OpenAI, bagian_sumber: str, img_bytes: bytes = None, label: str = "ekstraksi") -> list:
    """Satu panggilan ekstraksi untuk satu potongan teks ATAU satu gambar halaman."""
    prompt_text = f"""Dari bagian dokumen klaim berikut, ekstrak SELURUH baris item tagihan (billing) yang ada
di bagian ini saja (ini hanya potongan dari dokumen yang lebih panjang, jadi jangan menebak isi bagian lain).

PENTING: tabel tagihan RS biasanya berjudul "DETAIL LAYANAN", "RINCIAN BIAYA", "TINDAKAN", atau punya
kolom "Biaya Satuan"/"Total Biaya"/"Diskon". Jangan membuat item yang tidak ada. Kalau bagian ini tidak
berisi tabel tagihan, kembalikan {{"items": []}}.

JANGAN masukkan baris subtotal, total, atau judul kelompok, misalnya baris yang diawali "INSTALASI ...",
"Tagihan ...", "TOTAL", "SUBTOTAL", "JUMLAH". Hanya baris layanan/obat/alat/tindakan satuan.
harga_ditagih = nilai TOTAL baris itu (kolom jumlah/total biaya). harga_satuan = harga per satuan kalau
ada kolomnya, kalau tidak ada isi 0.

Jawab HANYA dengan satu JSON object, tanpa teks lain, format:
{{"items": [
  {{"kategori": "salah satu: Radiologi, Laboratorium, Obat, Alkes, Kamar, Dokter, Tindakan, Lainnya", "item": "nama item persis dari dokumen", "tanggal": "YYYY-MM-DD atau kosong kalau tidak ada", "quantity": 1, "harga_satuan": 0, "harga_ditagih": 0}}
]}}

{bagian_sumber}
"""
    content = [{"type": "text", "text": prompt_text}]
    if img_bytes is not None:
        b64 = base64.b64encode(img_bytes).decode()
        content.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}})

    resp = panggil_llm(
        client,
        label=label,
        model=MODEL_EKSTRAKSI,
        messages=[{"role": "user", "content": content}],
        response_format={"type": "json_object"},
        max_tokens=MAX_OUTPUT_EKSTRAKSI,
        temperature=0,
    )
    raw = teks_dari(resp)
    if resp.choices and resp.choices[0].finish_reason == "length":
        print(f"      [WARN] [{label}] output terpotong, sebagian item bagian ini mungkin hilang. "
              f"Kecilkan LLM_EKSTRAKSI_CHUNK_CHARS atau naikkan LLM_MAX_OUTPUT_EKSTRAKSI.")
    try:
        items = parse_json_response(raw)
    except (json.JSONDecodeError, IndexError) as e:
        print(f"      [WARN] [{label}] gagal parse JSON ({e}), mencoba selamatkan item parsial...")
        items = _salvage_json_items(raw)
        if not items:
            print(f"      [WARN] [{label}] tidak ada item yang dapat diselamatkan. Raw: {raw[:200]!r}")
            return None
        print(f"      [INFO] [{label}] berhasil menyelamatkan {len(items)} item dari output yang terpotong.")
    if isinstance(items, dict):
        items = items.get("items", [])
    return items if isinstance(items, list) else None


def ekstrak_item_tagihan(client: OpenAI, teks_klaim: str, gambar_halaman: list = None) -> list:
    """
    Ekstraksi dilakukan per potongan (teks dipecah per ~EKSTRAKSI_CHUNK_CHARS karakter,
    dan tiap gambar halaman dikirim sendiri-sendiri) supaya muat di context window server
    yang kecil. Hasil semua potongan digabung.
    """
    semua_item, gagal = [], 0

    potongan = _pecah_teks(teks_klaim, EKSTRAKSI_CHUNK_CHARS)
    for i, p in enumerate(potongan, 1):
        label = f"ekstraksi teks {i}/{len(potongan)}"
        hasil = _ekstrak_satu(client, f"DOKUMEN (teks, potongan {i} dari {len(potongan)}):\n{p}", label=label)
        if hasil is None:
            gagal += 1
        else:
            semua_item.extend(hasil)

    gambar_halaman = gambar_halaman or []
    for i, img in enumerate(gambar_halaman, 1):
        label = f"ekstraksi gambar {i}/{len(gambar_halaman)}"
        hasil = _ekstrak_satu(
            client,
            "SUMBER: gambar halaman yang dilampirkan (teksnya tidak terbaca otomatis, periksa secara visual dengan teliti).",
            img_bytes=img, label=label,
        )
        if hasil is None:
            gagal += 1
        else:
            semua_item.extend(hasil)

    if gagal:
        print(f"      [WARN] {gagal} potongan gagal diekstrak -- daftar item kemungkinan TIDAK LENGKAP, cek manual.")
    n_agregat = _tandai_agregat(semua_item)
    if n_agregat:
        print(f"      {n_agregat} baris subtotal/judul/harga-0 ditandai dan TIDAK ikut dianalisis sebagai item tagihan.")
    return semua_item


_POLA_AGREGAT = re.compile(r"^(instalasi|subtotal|sub\s*-?\s*total|grand\s+total)\b|^(jumlah|tagihan)\s*[:\-$]", re.I)
# "TOTAL" saja dianggap subtotal hanya kalau berdiri sendiri atau diikuti kata biaya/tagihan.
# Tanpa ini, tes lab seperti "Total Bilirubin" atau "Total Protein" ikut terbuang.
_POLA_TOTAL = re.compile(
    r"^total\b\s*[:\-]?\s*(?:$|(?:tagihan|biaya|bayar|pembayaran|harga|keseluruhan|rawat|layanan|instalasi|administrasi|akhir))",
    re.I)
_KATEGORI_VALID = {"radiologi", "laboratorium", "obat", "alkes", "kamar", "dokter", "tindakan", "lainnya"}


def _tandai_agregat(items: list) -> int:
    """
    Tandai baris subtotal murni ("INSTALASI ...", "TOTAL TAGIHAN") sebagai agregat.
    Item medis dengan harga 0 atau tidak terbaca TIDAK dibuang kecuali namanya jelas header kelompok,
    agar item tindakan/obat yang terbaca dari scan tetap muncul di laporan verifikasi.
    """
    n = 0
    for it in items:
        nama = str(it.get("item", "")).strip().lstrip("#").strip()
        kat = str(it.get("kategori", "")).strip().lower()
        if kat not in _KATEGORI_VALID:
            it["kategori"] = "" if "/" in kat or not kat else it.get("kategori", "")
        h = angka(it.get("harga_ditagih"))
        is_subtotal = bool(_POLA_AGREGAT.match(nama) or _POLA_TOTAL.match(nama) or kat.startswith(("instalasi", "tagihan")))
        if is_subtotal:
            it["_agregat"] = True
            it["_alasan_agregat"] = "subtotal/judul"
            n += 1
        elif h <= 0 and (len(nama) < 3 or is_subtotal):
            it["_agregat"] = True
            it["_alasan_agregat"] = "header/kosong"
            n += 1
    return n


# ---------------------------------------------------------------------------
# Tahap 2: Retrieval tarif pembanding dari database
# ---------------------------------------------------------------------------

def _muat_chunk_tarif(conn, kode_rs: str):
    """
    Ambil SEMUA chunk tarif milik satu RS lalu normalisasi vektornya, supaya kemiripan cosine
    bisa dihitung di Python (tanpa ekstensi pgvector). Berfungsi untuk kolom bertipe REAL[]
    maupun vector (pgvector mengembalikan teks '[0.1,0.2,...]', di-parse di bawah).
    Return: (meta, matriks) -- matriks None kalau RS ini belum punya chunk tarif.
    """
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT dc.isi_chunk, dc.no_chunk, d.nama_dokumen, dc.embedding
            FROM dokumen_chunk dc
            JOIN dokumen d ON d.id = dc.dokumen_id
            JOIN rumahsakit r ON r.id = d.rumah_sakit_id
            WHERE r.kode = %s;
            """,
            (kode_rs,),
        )
        rows = cur.fetchall()

    meta, vecs = [], []
    for isi, no, nama, emb in rows:
        if isinstance(emb, str):
            emb = json.loads(emb)
        meta.append({"isi_chunk": isi, "no_chunk": no, "nama_dokumen": nama})
        vecs.append(emb)
    if not vecs:
        return meta, None

    m = np.asarray(vecs, dtype=np.float32)
    m = m / (np.linalg.norm(m, axis=1, keepdims=True) + 1e-12)
    return meta, m


def _top_k_mirip(meta: list, matriks, query_vec, top_k: int) -> list:
    q = np.asarray(query_vec, dtype=np.float32)
    q = q / (np.linalg.norm(q) + 1e-12)
    skor = matriks @ q  # cosine similarity, karena kedua sisi sudah dinormalisasi
    idx = np.argsort(-skor)[:top_k]
    return [meta[int(i)] for i in idx]


def _tokens(teks) -> set:
    # Pertahankan token 1-2 karakter agar kode medis (K)/(TK) [dengan/tanpa kontras], AP/PA rontgen tidak hilang
    return {t for t in re.findall(r"[a-z0-9]+", str(teks).lower()) if t.isdigit() or len(t) >= 1}


def _bersihkan_nama_item_untuk_cari(nama: str) -> str:
    """Bersihkan nama dokter dan prefix '#' agar vector search mencocokkan tindakan medis, bukan nama dokter."""
    s = str(nama).strip().lstrip("#").strip()
    s = re.sub(r"[\(\[\-]?\s*dr\b\.?.*$", "", s, flags=re.I).strip()
    s = re.sub(r"\|\|.*$", "", s).strip()
    return s or str(nama).strip()


def _potong_relevan(chunk, nama_item: str, batas: int) -> str:
    """
    Kecilkan chunk ke baris-baris yang paling mirip dengan nama item (urutan asli dipertahankan),
    serta pertahankan baris header tabel (jika ada) agar nama kolom kelas tarif (Kelas 1/2/3) tidak hilang.
    """
    teks = re.sub(r"[ \t]+", " ", str(chunk)).strip()
    if len(teks) <= batas:
        return teks
    kt = _tokens(nama_item)
    baris = [b.strip() for b in teks.split("\n") if b.strip()]
    if not baris:
        return teks[:batas]

    # Cari baris header tabel (biasanya mengandung 'kelas', 'tarif', 'biaya', 'tindakan', atau '|')
    header_idx = set()
    for idx, b in enumerate(baris[:3]):
        b_low = b.lower()
        if any(h in b_low for h in ("kelas", "tarif", "biaya", "harga", "satuan", "layanan")) and "|" in b:
            header_idx.add(idx)

    urut = sorted(range(len(baris)), key=lambda i: (-(len(kt & _tokens(baris[i])) / (len(kt) or 1)), i))
    dipilih = set(header_idx)
    total = sum(len(baris[i]) + 1 for i in dipilih)

    for i in urut:
        if i in dipilih:
            continue
        skor = len(kt & _tokens(baris[i]))
        if skor == 0 and len(dipilih) > len(header_idx):
            break
        if total + len(baris[i]) + 1 > batas:
            continue
        dipilih.add(i)
        total += len(baris[i]) + 1

    if not dipilih:
        return teks[:batas]
    return "\n".join(baris[i] for i in sorted(dipilih))[:batas]


def _kandidat_tarif(meta: list, matriks, query_vec, nama_item: str, top_k: int, pool: int = 15) -> list:
    """Ambil `pool` kandidat terdekat (cosine), pangkas tiap chunk ke bagian relevan, rerank dengan kecocokan kata."""
    q = np.asarray(query_vec, dtype=np.float32)
    q = q / (np.linalg.norm(q) + 1e-12)
    skor = matriks @ q
    kt = _tokens(nama_item)
    kandidat, terlihat = [], set()
    for i in np.argsort(-skor)[:pool]:
        m = meta[int(i)]
        teks = _potong_relevan(m["isi_chunk"], nama_item, TARIF_CHUNK_CHARS)
        if teks in terlihat:
            continue
        terlihat.add(teks)
        leksikal = len(kt & _tokens(teks)) / len(kt) if kt else 0.0
        # Simpan isi_chunk_penuh untuk validasi angka di Python
        kandidat.append((float(skor[i]) + leksikal, {**m, "isi_chunk": teks, "isi_chunk_penuh": m["isi_chunk"]}))
    kandidat.sort(key=lambda x: -x[0])
    return [k for _, k in kandidat[:top_k]]


def cari_tarif_pembanding(
    conn, embed_model: SentenceTransformer, kode_rs: str, nama_item: str, top_k: int = TOP_K_TARIF, data=None
) -> list:
    """Cari chunk tarif paling mirip untuk satu item, dibatasi ke rumah sakit yang sesuai."""
    meta, matriks = data if data is not None else _muat_chunk_tarif(conn, kode_rs)
    if matriks is None:
        return []
    nama_cari = _bersihkan_nama_item_untuk_cari(nama_item)
    return _kandidat_tarif(meta, matriks, embed_model.encode([nama_cari])[0], nama_cari, top_k)


def lengkapi_item_dengan_tarif(
    conn, embed_model: SentenceTransformer, kode_rs: str, items: list
) -> list:
    # Chunk tarif dimuat SEKALI per klaim, bukan per item.
    data = _muat_chunk_tarif(conn, kode_rs)
    if data[1] is None:
        print(f"      [WARN] Tidak ada chunk tarif untuk kode_rs '{kode_rs}' di database. "
              f"Semua item akan tanpa tarif pembanding. Pastikan dokumen tarif sudah di-ingest dengan kode RS yang sama.")
    else:
        panjang = sorted(len(m["isi_chunk"]) for m in data[0])
        print(f"      {len(data[0])} chunk tarif dimuat untuk '{kode_rs}' "
              f"(panjang median {panjang[len(panjang) // 2]}, terpanjang {panjang[-1]} karakter)")
        if panjang[-1] > 2000:
            print("      [INFO] Ada chunk tarif sepanjang halaman; dipangkas ke baris yang relevan per item saat dikirim ke LLM.")

    nama_cari_semua = [_bersihkan_nama_item_untuk_cari(it.get("item", "")) for it in items]
    vektor = embed_model.encode([n or " " for n in nama_cari_semua]) if items else []
    for item, nama_cari, v in zip(items, nama_cari_semua, vektor):
        if not nama_cari or data[1] is None or item.get("_agregat"):
            item["tarif_pembanding"] = []
        else:
            item["tarif_pembanding"] = _kandidat_tarif(data[0], data[1], v, nama_cari, TOP_K_TARIF)
    return items


# ---------------------------------------------------------------------------
# Tahap 3: Analisis akhir
# ---------------------------------------------------------------------------

_KUNCI_ITEM_PROMPT = ("kategori", "item", "tanggal", "quantity", "harga_satuan", "harga_ditagih")


def _ringkas_item_untuk_prompt(items: list, batas_chunk: int, top_k: int) -> list:
    """Hanya field yang dibutuhkan model; tiap chunk tarif dipangkas ke baris relevan (bukan satu halaman penuh)."""
    hasil = []
    for it in items:
        d = {k: it.get(k) for k in _KUNCI_ITEM_PROMPT if it.get(k) not in (None, "")}
        d["tarif_pembanding"] = [
            {"nama_dokumen": t.get("nama_dokumen"), "no_chunk": t.get("no_chunk"),
             "isi_chunk": _potong_relevan(t.get("isi_chunk", ""), it.get("item", ""), batas_chunk)}
            for t in (it.get("tarif_pembanding") or [])[:top_k]
        ]
        hasil.append(d)
    return hasil


def _susun_prompt_penuh(system_prompt: str, teks_klaim: str, items: list, batas_chunk: int = None, top_k: int = TOP_K_TARIF) -> str:
    ringkas = _ringkas_item_untuk_prompt(items, batas_chunk or TARIF_CHUNK_CHARS, top_k)
    # Batasi teks_klaim agar tidak meledakkan token context window server
    teks_klaim_aman = teks_klaim[:25000] if len(teks_klaim) > 25000 else teks_klaim
    return f"""{system_prompt}

---

# DOKUMEN KLAIM (INPUT)

{teks_klaim_aman}

---

# ITEM TAGIHAN + HASIL PENCARIAN TARIF (dari database, lihat aturan penggunaan di atas)

WAJIB: field "tabel_item" pada output JSON harus berisi PERSIS {len(items)} baris, satu baris untuk
SETIAP item di bawah ini, tanpa terkecuali. Jangan meringkas, menggabungkan, mengelompokkan, atau
menghilangkan item apa pun walau isinya mirip atau kelihatan berulang -- itu tetap baris billing yang
terpisah dan harus tetap terlihat terpisah di laporan.

{json.dumps(ringkas, ensure_ascii=False, separators=(",", ":"))}
"""


def _laporkan_ukuran(system_prompt: str, teks_klaim: str, items: list) -> dict:
    """Rincian ukuran prompt per komponen, untuk melihat BAGIAN MANA yang membengkak."""
    mentah = sum(len(str(t.get("isi_chunk", ""))) for it in items for t in it.get("tarif_pembanding", []))
    ringkas = _ringkas_item_untuk_prompt(items, TARIF_CHUNK_CHARS, TOP_K_TARIF)
    setelah = sum(len(t["isi_chunk"]) for d in ringkas for t in d["tarif_pembanding"])
    json_items = len(json.dumps(ringkas, ensure_ascii=False, separators=(",", ":")))
    print("      [UKURAN PROMPT] komponen analisis penuh (karakter, ~token pada 2,5 kar/token):")
    for nama, n in (("system prompt", len(system_prompt)), ("teks klaim", len(teks_klaim)),
                    ("tarif pembanding MENTAH (sebelum dipangkas)", mentah),
                    ("tarif pembanding setelah dipangkas", setelah), ("JSON item + tarif (dikirim)", json_items)):
        print(f"        - {nama}: {n:,} (~{_taksir_token('x' * n):,} token)")
    besar = sorted(items, key=lambda it: -sum(len(str(t.get("isi_chunk", ""))) for t in it.get("tarif_pembanding", [])))[:3]
    for it in besar:
        print(f"        - item termahal: {str(it.get('item'))[:40]!r} -> "
              f"{sum(len(str(t.get('isi_chunk', ''))) for t in it.get('tarif_pembanding', [])):,} karakter tarif mentah")
    return {"system": len(system_prompt), "teks_klaim": len(teks_klaim), "tarif_mentah": mentah, "tarif_dipangkas": setelah}


def _analisis_klaim_penuh(
    client: OpenAI, system_prompt: str, teks_klaim: str, items_dengan_tarif: list, prompt: str = None
) -> dict:
    jumlah_item = len(items_dengan_tarif)
    prompt = prompt or _susun_prompt_penuh(system_prompt, teks_klaim, items_dengan_tarif)

    resp = panggil_llm(
        client,
        label="analisis",
        model=MODEL_ANALISIS,
        messages=[{"role": "user", "content": prompt}],
        response_format={"type": "json_object"},
        max_tokens=MAX_OUTPUT_ANALISIS,
        temperature=0,
    )
    raw = teks_dari(resp)
    try:
        hasil = parse_json_response(raw)
    except (json.JSONDecodeError, IndexError) as e:
        sys.exit(f"Gagal parse hasil analisis dari LLM: {e}\nRaw: {raw[:1000]}")

    if isinstance(hasil, list):
        # Kadang model cuma balikin array tabel_item-nya saja, lupa bungkus jadi objek lengkap.
        print("      [WARN] LLM mengembalikan list, bukan objek lengkap. Membungkusnya otomatis -- "
              "narasi/keputusan akhir kemungkinan tidak lengkap, cek manual.")
        hasil = {"tabel_item": hasil}

    if not isinstance(hasil, dict):
        sys.exit(f"Hasil analisis dari LLM bentuknya tidak dikenali (tipe: {type(hasil)}).\nRaw: {raw[:1000]}")

    jumlah_hasil = len(hasil.get("tabel_item", []))
    if jumlah_hasil != jumlah_item:
        print(
            f"      [WARN] Item hilang saat analisis akhir: {jumlah_item} item dikirim, "
            f"hanya {jumlah_hasil} yang muncul di tabel_item. Merekonstruksi item yang terpotong..."
        )
        item_ada = {str(r.get("item", "")).strip().lower() for r in (hasil.get("tabel_item") or []) if isinstance(r, dict)}
        for it in items_dengan_tarif:
            kunci = str(it.get("item", "")).strip().lower()
            if kunci not in item_ada:
                hasil.setdefault("tabel_item", []).append({
                    "kategori": it.get("kategori", ""),
                    "item": it.get("item", ""),
                    "tanggal": it.get("tanggal", ""),
                    "quantity": it.get("quantity", 1),
                    "harga_satuan": it.get("harga_satuan", 0),
                    "harga_ditagih": it.get("harga_ditagih", 0),
                    "tarif_rs": None,
                    "selisih": None,
                    "sumber_tarif": "",
                    "keterangan": "TIDAK ADA TARIF PEMBANDING (perlu verifikasi manual)",
                })

    return hasil


# ---------------------------------------------------------------------------
# Tahap 3 (alternatif): analisis berjenjang untuk context window kecil
#   a) item dianalisis per batch kecil (hanya bandingkan harga vs tarif)
#   b) angka total/selisih dihitung di Python, bukan oleh LLM
#   c) satu panggilan terakhir menyusun narasi & keputusan dari ringkasan angka
# ---------------------------------------------------------------------------

def _taksir_token(teks: str) -> int:
    return int(len(teks) / CHARS_PER_TOKEN)


def _potong_token(teks: str, token: int) -> str:
    return teks[: max(int(token * CHARS_PER_TOKEN), 0)]


def _analisis_satu_batch(client: OpenAI, batch: list, label: str):
    """Bandingkan harga vs tarif untuk satu batch kecil. Return list baris, atau None kalau gagal."""
    ringkas = []
    for it in batch:
        ringkas.append({
            "item": it.get("item"),
            "harga_ditagih": it.get("harga_ditagih"),
            "tarif_pembanding": [
                {"nama_dokumen": t.get("nama_dokumen"), "isi_chunk": _potong_relevan(t.get("isi_chunk", ""), it.get("item", ""), TARIF_CHUNK_CHARS)}
                for t in it.get("tarif_pembanding", [])
            ],
        })
    prompt = f"""Kamu verifikator klaim rumah sakit. Ada {len(batch)} item tagihan di bawah. Untuk SETIAP item
(tepat {len(batch)} baris, urutan sama persis dengan input), bandingkan harga_ditagih dengan tarif_pembanding
(dari database tarif RS yang sama).

Aturan:
- tarif_rs: angka murni (tanpa "Rp", tanpa titik/koma ribuan) dari kandidat yang JELAS merujuk layanan/item yang sama.
  Kalau tidak ada kandidat yang jelas cocok, isi null. Jangan mengarang angka.
- sumber_tarif: nama_dokumen kandidat yang dipakai, atau "" kalau tarif_rs null.
- keterangan: salah satu persis: "SESUAI", "OVER TARIF", "UNDER TARIF", "TIDAK ADA TARIF PEMBANDING".

Jawab HANYA satu JSON object: {{"tabel_item": [{{"item": "...", "tarif_rs": 12500, "sumber_tarif": "...", "keterangan": "SESUAI"}}]}}

INPUT:
{json.dumps(ringkas, ensure_ascii=False)}
"""
    try:
        resp = panggil_llm(
            client, label=label, model=MODEL_ANALISIS,
            messages=[{"role": "user", "content": prompt}],
            response_format={"type": "json_object"},
            max_tokens=MAX_OUTPUT_BATCH, temperature=0,
        )
        data = parse_json_response(teks_dari(resp))
    except Exception as e:
        print(f"      [WARN] [{label}] gagal ({e.__class__.__name__}: {str(e)[:200]}); batch ini ditandai cek manual.")
        return None
    rows = data.get("tabel_item") if isinstance(data, dict) else data
    return rows if isinstance(rows, list) else None


def _gabung_baris(batch: list, rows) -> list:
    """Kolom identitas item diambil dari hasil ekstraksi (bukan dari LLM) supaya tidak berubah/hilang."""
    by_name = {}
    for r in (rows or []):
        if isinstance(r, dict):
            by_name.setdefault(str(r.get("item", "")).strip().lower(), r)
    sama_panjang = rows is not None and len(rows) == len(batch)

    hasil = []
    for i, it in enumerate(batch):
        r = rows[i] if sama_panjang and isinstance(rows[i], dict) else by_name.get(str(it.get("item", "")).strip().lower())
        baris = {k: it.get(k) for k in ("kategori", "item", "tanggal", "quantity", "harga_satuan", "harga_ditagih")}
        if r:
            baris["tarif_rs"] = r.get("tarif_rs")
            baris["sumber_tarif"] = r.get("sumber_tarif", "")
            baris["keterangan"] = r.get("keterangan", "")
        else:
            baris["tarif_rs"] = None
            baris["sumber_tarif"] = ""
            baris["keterangan"] = "Tidak teranalisis (cek manual)"
        hasil.append(baris)
    return hasil


def ekstrak_identitas(client: OpenAI, teks_klaim: str, gambar_halaman: list = None) -> dict:
    """Identitas pasien diekstrak dari teks klaim atau gambar cover/scan pertama kalau teks minim."""
    bagian = [b for b in re.split(r"(?==== DOKUMEN: )", teks_klaim) if b.strip()]
    batas_total = 12000
    per_dokumen = max(batas_total // max(len(bagian), 1), 1500)
    kutipan = "".join(b[:per_dokumen] for b in bagian)[:batas_total]
    prompt = f"""Dari kutipan dokumen klaim rumah sakit berikut, ambil identitas pasien dan data rawat.
Isi HANYA yang tertulis jelas di dokumen. Kalau tidak ada, isi "" (kosong). Jangan menebak.
Tanggal dalam format YYYY-MM-DD.

Jawab HANYA satu JSON object:
{{"nama_pasien": "", "no_peserta": "", "rumah_sakit": "", "tgl_masuk": "", "tgl_keluar": "", "program": "", "diagnosa": ""}}

KUTIPAN:
{kutipan}
"""
    content = [{"type": "text", "text": prompt}]
    # Jika teks klaim sangat sedikit dan ada gambar scan (halaman pertama biasanya resume medis / form pengajuan)
    if len(kutipan.strip()) < 300 and gambar_halaman:
        b64 = base64.b64encode(gambar_halaman[0]).decode()
        content = [
            {"type": "text", "text": "Dari gambar dokumen klaim ini (halaman depan/resume medis), ambil identitas pasien dan data rawat. Jawab HANYA JSON object dengan format: {\"nama_pasien\": \"\", \"no_peserta\": \"\", \"rumah_sakit\": \"\", \"tgl_masuk\": \"\", \"tgl_keluar\": \"\", \"program\": \"\", \"diagnosa\": \"\"}"},
            {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}}
        ]

    idn = {}
    try:
        resp = panggil_llm(
            client, label="identitas", model=MODEL_ANALISIS,
            messages=[{"role": "user", "content": content}],
            response_format={"type": "json_object"}, max_tokens=600, temperature=0,
        )
        data = parse_json_response(teks_dari(resp))
        idn = data if isinstance(data, dict) else {}
    except Exception as e:
        print(f"      [WARN] Ekstraksi identitas gagal ({e.__class__.__name__}: {str(e)[:150]}); identitas dikosongkan.")

    hasil = {k: str(idn.get(k, "") or "") for k in ("nama_pasien", "no_peserta", "rumah_sakit", "program", "diagnosa")}
    masuk, keluar = str(idn.get("tgl_masuk", "") or ""), str(idn.get("tgl_keluar", "") or "")
    hasil["periode_rawat"] = f"{masuk} s.d. {keluar}" if masuk and keluar else (masuk or keluar)
    try:
        # LOS dihitung dari tanggal (tgl keluar - tgl masuk), bukan dari tebakan model
        hasil["los_hari"] = (datetime.strptime(keluar, "%Y-%m-%d") - datetime.strptime(masuk, "%Y-%m-%d")).days
    except ValueError:
        hasil["los_hari"] = ""
    hasil["total_tagihan_rs"] = 0  # diisi dari baris "Tagihan ..." di validasi_hasil
    return hasil


def _analisis_klaim_berjenjang(client: OpenAI, system_prompt: str, teks_klaim: str, items: list, gambar_halaman: list = None) -> dict:
    """
    Jalur untuk context kecil. LLM hanya dipakai untuk (a) mencocokkan tarif per batch kecil dan
    (b) mengekstrak identitas. Angka, status, dan rekap dihitung di Python (validasi_hasil).
    system_prompt tidak dipakai di jalur ini (terlalu panjang untuk model kecil); parameternya
    dipertahankan supaya signature sama dengan jalur penuh.
    """
    baris_tabel = []
    n_batch = (len(items) + ANALISIS_BATCH_ITEM - 1) // ANALISIS_BATCH_ITEM
    for b in range(n_batch):
        batch = items[b * ANALISIS_BATCH_ITEM:(b + 1) * ANALISIS_BATCH_ITEM]
        rows = _analisis_satu_batch(client, batch, f"analisis batch {b + 1}/{n_batch}")
        baris_tabel.extend(_gabung_baris(batch, rows))
    return {"identitas": ekstrak_identitas(client, teks_klaim, gambar_halaman), "tabel_item": baris_tabel}


# ---------------------------------------------------------------------------
# Validasi & perhitungan deterministik (di Python, bukan oleh LLM)
# ---------------------------------------------------------------------------

def _norm_nama(s) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def _angka_dalam_teks(teks: str) -> set:
    """Semua bilangan bulat yang tertulis di teks (dengan/tanpa pemisah ribuan dan desimal ,00)."""
    hasil = set()
    for tok in re.findall(r"\d[\d.,]*", teks or ""):
        tok = tok.rstrip(".,")
        varian = {tok}
        m = re.search(r"[.,]\d{1,2}$", tok)
        if m:
            varian.add(tok[:m.start()])
        for v in varian:
            d = re.sub(r"\D", "", v)
            if d:
                hasil.add(int(d))
    return hasil


def _bandingkan(harga, qty, tarif):
    """
    Rumus dari system prompt: tarif RS adalah tarif PER UNIT. Kalau quantity > 1, tarif seharusnya =
    quantity x tarif; selisih = total ditagih - tarif seharusnya. harga_ditagih dianggap total baris.
    Return (keterangan, selisih, persen_selisih, tarif_seharusnya).
    """
    tarif = float(tarif)
    harga = float(angka(harga))
    q = angka(qty, 1) or 1
    seharusnya = tarif * q if q > 1 else tarif
    selisih = harga - seharusnya
    persen = (selisih / seharusnya * 100) if seharusnya else 0.0
    if abs(selisih) <= max(seharusnya * 0.01, 1):
        return "SESUAI", selisih, persen, seharusnya
    return ("OVER TARIF" if selisih > 0 else "UNDER TARIF"), selisih, persen, seharusnya


# Klasifikasi selain soal tarif yang mungkin diberikan model (jalur penuh). Tidak ditimpa oleh perhitungan tarif.
_KELAS_LAIN = re.compile(r"double|duplicate|unbundling|overutilization|indikasi medis|klarifikasi", re.I)


def validasi_hasil(hasil: dict, items: list = None, items_agregat: list = None, hitung_ulang: bool = False) -> dict:
    """
    Pemeriksaan & koreksi deterministik atas keluaran LLM:
      - tarif 0 / tarif yang angkanya tidak ada di teks tarif pembanding dianggap tidak ada pembanding
      - keterangan tarif (SESUAI / OVER TARIF / UNDER TARIF) dihitung dari angka, bukan dipercaya dari model;
        klasifikasi lain (DOUBLE BILLING, UNBUNDLING, dst) dari model dipertahankan
      - rekonsiliasi jumlah item vs total tagihan, deteksi nama item ganda
      - hitung_ulang=True (jalur berjenjang): status, angka keputusan, rekap, confidence dihitung di sini
    """
    peringatan = []
    rows = [r for r in (hasil.get("tabel_item") or []) if isinstance(r, dict)]
    n = len(rows)
    identitas = hasil.get("identitas") if isinstance(hasil.get("identitas"), dict) else {}
    hasil["identitas"] = identitas
    kp = hasil.get("keputusan_akhir") if isinstance(hasil.get("keputusan_akhir"), dict) else {}
    status_awal = str(kp.get("status", "")).strip().upper()

    teks_kandidat = {}
    for it in (items or []):
        kunci = _norm_nama(it.get("item", ""))
        gabung = " ".join(str(t.get("isi_chunk_penuh", t.get("isi_chunk", ""))) for t in it.get("tarif_pembanding", []))
        teks_kandidat.setdefault(kunci, gabung)

    # 1) tarif & keterangan
    nol = tak_berdasar = ket_dikoreksi = 0
    for r in rows:
        mentah = r.get("tarif_rs")
        t = angka(mentah, None) if mentah not in (None, "") else None
        if t is not None and t <= 0:
            nol += 1
            t = None
        if t is not None:
            kunci = _norm_nama(r.get("item", ""))
            if kunci in teks_kandidat and int(round(t)) not in _angka_dalam_teks(teks_kandidat[kunci]):
                tak_berdasar += 1
                t = None
        asli = str(r.get("keterangan", ""))
        kelas_lain = bool(_KELAS_LAIN.search(asli))
        r["sumber_tarif_klaim_model"] = r.get("sumber_tarif", "")
        if t is None:
            r["tarif_rs"], r["selisih"], r["sumber_tarif"] = None, None, ""
            r["persen_selisih"], r["tarif_seharusnya"] = None, None
            if not kelas_lain:
                r["keterangan"] = "TIDAK ADA TARIF PEMBANDING"
        else:
            ket, sel_, persen, seh = _bandingkan(r.get("harga_ditagih"), r.get("quantity"), t)
            if "sesuai" in asli.lower() and ket != "SESUAI" and not kelas_lain:
                ket_dikoreksi += 1
            r["tarif_rs"], r["selisih"] = t, sel_
            r["persen_selisih"], r["tarif_seharusnya"] = round(persen, 1), seh
            if not kelas_lain:
                r["keterangan"] = ket
    tanpa = sum(1 for r in rows if r.get("tarif_rs") is None)
    cakupan = (n - tanpa) / n if n else 0

    if nol:
        peringatan.append(f"{nol} item diberi tarif 0 oleh model; dianggap TIDAK ADA TARIF PEMBANDING (bukan selisih sebesar harga).")
    if tak_berdasar:
        peringatan.append(f"{tak_berdasar} item: angka tarif dari model tidak tertulis di teks tarif pembanding "
                          f"(kemungkinan dikarang atau disalin dari harga tagihan); dianggap TIDAK ADA TARIF PEMBANDING.")
    if ket_dikoreksi:
        peringatan.append(f"{ket_dikoreksi} item ditandai 'SESUAI' oleh model padahal selisihnya tidak nol; "
                          f"keterangan diganti hasil perhitungan.")
    if n and tanpa:
        peringatan.append(f"Tarif pembanding tidak ada untuk {tanpa} dari {n} item ({tanpa / n:.0%}). "
                          f"BELUM DAPAT DIVERIFIKASI TERHADAP TARIF RS (bukan 'sesuai').")

    # 2) rekonsiliasi total
    total_rs = angka(identitas.get("total_tagihan_rs"))
    if not total_rs and items_agregat:
        kandidat_total = [angka(a.get("harga_ditagih")) for a in items_agregat if _norm_nama(a.get("item", "")).startswith("tagihan")]
        total_rs = max(kandidat_total) if kandidat_total else 0
        if total_rs:
            identitas["total_tagihan_rs"] = total_rs
    jml_harga = sum(angka(r.get("harga_ditagih")) for r in rows)
    jml_x_qty = sum(angka(r.get("harga_ditagih")) * (angka(r.get("quantity"), 1) or 1) for r in rows)
    if total_rs > 0 and rows and not any(abs(x - total_rs) / total_rs <= 0.01 for x in (jml_harga, jml_x_qty)):
        peringatan.append(f"Jumlah item ({rupiah(jml_harga)}; dikali qty {rupiah(jml_x_qty)}) tidak cocok dengan total tagihan RS "
                          f"({rupiah(total_rs)}). Kemungkinan: item ganda dari beberapa dokumen, item belum lengkap, atau kolom harga salah baca.")
    if items_agregat:
        peringatan.append(f"{len(items_agregat)} baris subtotal/judul/harga-0 dari dokumen diabaikan dari perhitungan (tersimpan di baris_agregat_diabaikan).")
        hasil["baris_agregat_diabaikan"] = [{"item": a.get("item"), "harga": a.get("harga_ditagih"), "alasan": a.get("_alasan_agregat")} for a in items_agregat]

    # 3) nama item ganda
    hitung = collections.Counter(_norm_nama(str(r.get("item", "")).lstrip("#")) for r in rows)
    ganda = {k for k, v in hitung.items() if k and v > 1}
    nilai_ganda = sum(angka(r.get("harga_ditagih")) for r in rows if _norm_nama(str(r.get("item", "")).lstrip("#")) in ganda)
    n_ganda = sum(hitung[k] for k in ganda)
    if ganda:
        peringatan.append(f"{n_ganda} baris memakai {len(ganda)} nama item yang muncul lebih dari sekali "
                          f"(bisa dari dokumen berbeda atau tanggal berbeda), nilai {rupiah(nilai_ganda)}. "
                          f"INDIKASI/ANOMALI YANG MEMERLUKAN VERIFIKASI LANJUT; total bisa terhitung dua kali.")

    if hitung_ulang:
        total = jml_harga
        over = [r for r in rows if r.get("keterangan") == "OVER TARIF"]
        koreksi = sum(r["selisih"] for r in over)
        klarif = sum(angka(r.get("harga_ditagih")) for r in rows if r.get("tarif_rs") is None)
        rekom = max(total - koreksi - klarif, 0)
        # Mode ini hanya memverifikasi TARIF. Medical necessity, LOS, JKK/PLKK, unbundling, dst tidak diperiksa per item,
        # jadi sistem tidak boleh menyatakan layak (prinsip utama prompt: jangan menyimpulkan LAYAK hanya karena data terlihat wajar).
        status = "PERLU KLARIFIKASI" if (n and cakupan < 0.5) else "PERLU REVIEW MEDIS LANJUT"

        def sev(rasio, tinggi, sedang):
            return "High" if rasio > tinggi else ("Medium" if rasio > sedang else "Low")

        hasil["keputusan_akhir"] = {
            "status": status, "total_diajukan": total, "koreksi_terverifikasi": koreksi,
            "perlu_klarifikasi": klarif, "rekomendasi_pembayaran": rekom,
            "alasan": [
                f"{len(over)} item OVER TARIF terhadap tarif pembanding (potensi koreksi {rupiah(koreksi)}).",
                f"{tanpa} dari {n} item tidak punya tarif pembanding yang terverifikasi ({rupiah(klarif)}); cakupan tarif {cakupan:.0%}.",
                "Hanya verifikasi tarif yang dijalankan sistem pada mode ini; medical necessity, LOS, dan JKK/PLKK belum diverifikasi.",
                "Status dihitung otomatis dan bukan keputusan akhir; wajib ditinjau verifikator manusia.",
            ],
        }
        hasil["rekap_temuan"] = [
            {"jenis_temuan": "Overcharge (OVER TARIF)", "jumlah_item": len(over), "nilai": koreksi,
             "severity": sev(koreksi / total if total else 0, 0.05, 0),
             "rekomendasi": "Mintakan penjelasan/koreksi ke RS untuk item OVER TARIF."},
            {"jenis_temuan": "TIDAK ADA TARIF PEMBANDING", "jumlah_item": tanpa, "nilai": klarif,
             "severity": sev(klarif / total if total else 0, 0.5, 0.1),
             "rekomendasi": "Cek manual ke dokumen tarif/PKS; periksa apakah dokumen tarif RS sudah di-ingest dengan kode RS yang benar."},
            {"jenis_temuan": "Double/Duplicate Billing (nama item sama)", "jumlah_item": n_ganda, "nilai": nilai_ganda,
             "severity": "Medium" if ganda else "Low",
             "rekomendasi": ("INDIKASI/ANOMALI YANG MEMERLUKAN VERIFIKASI LANJUT: bandingkan tanggal dan dokumen sumber."
                             if ganda else "Tidak ditemukan nama item ganda berdasarkan data yang tersedia (bukan berarti dipastikan tidak ada).")},
        ]
        hasil["confidence_score"] = round(min(70, 20 + 50 * cakupan))
        peringatan.append("Mode berjenjang: hanya verifikasi tarif yang dijalankan terstruktur. Unbundling, overutilization, medical necessity, "
                          "date mismatch, quantity anomaly, LOS, dan JKK/PLKK TIDAK diperiksa per item; wajib ditinjau manual.")
    else:
        # jalur penuh: hasil model dipertahankan, hanya dicek kewajarannya
        if n and cakupan < 0.5:
            if status_awal.startswith("LAYAK"):   # startswith: "TIDAK LAYAK" tidak boleh ikut ditimpa
                kp["status_model_asli"] = kp.get("status")
                kp["status"] = "PERLU KLARIFIKASI"
                hasil["keputusan_akhir"] = kp
                peringatan.append(f"Status dari model ('{kp['status_model_asli']}') diubah otomatis menjadi PERLU KLARIFIKASI karena kurang dari 50% item punya tarif pembanding.")
            conf = hasil.get("confidence_score")
            if isinstance(conf, (int, float)) and conf > 50:
                hasil["confidence_score_model_asli"] = conf
                hasil["confidence_score"] = 50
                peringatan.append(f"Confidence score diturunkan dari {conf} ke 50 (cakupan tarif {cakupan:.0%}).")
            if any("sesuai" in str(a).lower() and "tarif" in str(a).lower() for a in (kp.get("alasan") or [])):
                peringatan.append("Salah satu 'alasan' menyebut tagihan sesuai tarif, padahal tarif belum terverifikasi. Abaikan alasan itu.")
        if "KOREKSI" in status_awal and angka(kp.get("koreksi_terverifikasi")) == 0:
            peringatan.append("Status menyebut 'koreksi' tetapi koreksi_terverifikasi = 0. Status dan angka tidak konsisten.")

    hasil["peringatan"] = peringatan
    hasil["cakupan_tarif"] = round(cakupan, 3)
    return hasil


_FRASA_BOCOR = ("dihitung sistem", "tidak boleh dihitung", "jangan dihitung", "jawab hanya", "json object", "teman")
_KATA_DOK_MEDIS = re.compile(r"resume|medis|plkk|kk[123]|igd|kronologi|lembur|saksi|surat|skp|rekam", re.I)


def _narasi_keuangan(hasil: dict) -> str:
    """Bagian narasi tentang tagihan dan tarif, disusun dari angka yang sudah dihitung (tanpa LLM)."""
    i, k = hasil.get("identitas", {}), hasil.get("keputusan_akhir", {})
    return (
        f"**Identitas:** {i.get('nama_pasien') or '-'} | {i.get('rumah_sakit') or '-'} | periode rawat {i.get('periode_rawat') or '-'}"
        f" | LOS {i.get('los_hari') if i.get('los_hari') not in (None, '') else '-'} hari.\n\n"
        f"**Tagihan:** total menurut dokumen {rupiah(i.get('total_tagihan_rs'))}; jumlah item yang dianalisis {rupiah(k.get('total_diajukan'))}.\n\n"
        f"- OVER TARIF: potensi koreksi {rupiah(k.get('koreksi_terverifikasi'))}.\n"
        f"- TIDAK ADA TARIF PEMBANDING (belum terverifikasi): {rupiah(k.get('perlu_klarifikasi'))}.\n"
        f"- Cakupan tarif {hasil.get('cakupan_tarif', 0):.0%}. Status otomatis: {k.get('status', '')}.\n\n"
        f"Verifikator perlu memeriksa item berstatus OVER TARIF dan TIDAK ADA TARIF PEMBANDING, serta peringatan di atas."
    )


def _kutipan_medis(teks_klaim: str, batas_chars: int) -> str:
    """Ambil awal dokumen yang tampak medis/PLKK lebih dulu (resume, KK1-3, IGD, surat, dst)."""
    bagian = [b for b in re.split(r"(?==== DOKUMEN: )", teks_klaim) if b.strip()]
    bagian.sort(key=lambda b: 0 if _KATA_DOK_MEDIS.search(b.split("\n", 2)[1] if "\n" in b else b[:100]) else 1)
    per = max(batas_chars // max(len(bagian), 1), 1200)
    return "".join(b[:per] for b in bagian)[:batas_chars]


def _analisis_medis_jkk(client: OpenAI, system_prompt: str, teks_klaim: str):
    """
    Bagian B, C (kelengkapan, medical necessity, LOS) dan D (JKK/PLKK) dari system prompt, ditulis model
    dari kutipan dokumen medis. Return teks Markdown, atau None kalau context tidak cukup/keluaran janggal.
    """
    sisa = LLM_MAX_CONTEXT - MAX_OUTPUT_FINAL - 900 - _taksir_token(system_prompt)
    if sisa < 1500:
        print(f"      [WARN] Context {LLM_MAX_CONTEXT} terlalu kecil untuk system prompt + dokumen medis; "
              f"bagian medical necessity/LOS/JKK tidak dianalisis otomatis.")
        return None
    kutipan = _kutipan_medis(teks_klaim, min(int(sisa * CHARS_PER_TOKEN), 12000))
    prompt = f"""{system_prompt}

---

# INSTRUKSI KHUSUS UNTUK PANGGILAN INI (menggantikan format output JSON di atas)

Kerjakan HANYA bagian B (input data klaim), C (kelengkapan berkas, medical necessity, LOS), dan D (verifikasi JKK/PLKK)
dari proses verifikasi di atas, berdasarkan KUTIPAN DOKUMEN di bawah. Pisahkan fakta dokumen, analisis, dan data yang
belum tersedia. Kalau bukti tidak ada di kutipan, tulis: "TIDAK DAPAT DIVERIFIKASI — data/dokumen pendukung tidak tersedia."
Jangan membahas tarif atau tagihan di sini, dan jangan mengarang diagnosis, tanggal, atau kronologi.

Jawab dengan satu JSON object: {{"narasi_medis_jkk": "teks Markdown"}}

# KUTIPAN DOKUMEN

{kutipan}
"""
    try:
        resp = panggil_llm(
            client, label="analisis medis & JKK", model=MODEL_ANALISIS,
            messages=[{"role": "user", "content": prompt}],
            response_format={"type": "json_object"}, max_tokens=MAX_OUTPUT_FINAL, temperature=0,
        )
        data = parse_json_response(teks_dari(resp))
    except Exception as e:
        print(f"      [WARN] Analisis medis/JKK gagal ({e.__class__.__name__}: {str(e)[:150]}).")
        return None
    teks = None
    if isinstance(data, dict):
        teks = data.get("narasi_medis_jkk") or data.get("narasi_lengkap")
    if isinstance(teks, str) and len(teks) >= 200 and not any(f in teks.lower() for f in _FRASA_BOCOR):
        return teks
    print("      [WARN] Keluaran analisis medis/JKK tidak layak (terlalu pendek atau menyalin instruksi); dibuang.")
    return None


def _tulis_narasi(client: OpenAI, hasil: dict, system_prompt: str = "", teks_klaim: str = "") -> None:
    """Narasi = bagian keuangan (deterministik) + bagian medis/JKK (draft model, kalau muat dan layak)."""
    bagian = ["## Ringkasan Verifikasi Tagihan\n", _narasi_keuangan(hasil), "\n\n## Analisis Medis & JKK\n"]
    medis = _analisis_medis_jkk(client, system_prompt, teks_klaim) if system_prompt else None
    if medis:
        bagian.append("> DRAFT dari model bahasa. Wajib ditinjau dan dikoreksi verifikator manusia.\n\n" + medis)
    else:
        bagian.append("BELUM DAPAT DIVERIFIKASI secara otomatis (kelengkapan berkas, medical necessity, LOS, hubungan kerja/PLKK). "
                      "Wajib ditinjau manual oleh verifikator.")
    hasil["narasi_lengkap"] = "\n".join(bagian)


def analisis_klaim(
    client: OpenAI, system_prompt: str, teks_klaim: str, items_dengan_tarif: list, gambar_halaman: list = None
) -> dict:
    """
    Pilih analisis penuh (1 panggilan) kalau prompt SEBENARNYA muat, selain itu berjenjang. Keputusan memakai
    prompt yang benar-benar disusun (bukan perkiraan terpisah yang bisa meleset karena format JSON berbeda).
    """
    items_agregat = [i for i in items_dengan_tarif if i.get("_agregat")]
    items = [i for i in items_dengan_tarif if not i.get("_agregat")]

    _laporkan_ukuran(system_prompt, teks_klaim, items)
    batas_prompt = int(LLM_MAX_CONTEXT * 0.9) - MAX_OUTPUT_ANALISIS
    if BATAS_PROMPT_TOKEN:
        batas_prompt = min(batas_prompt, BATAS_PROMPT_TOKEN)

    prompt = _susun_prompt_penuh(system_prompt, teks_klaim, items)
    tok = _taksir_token(prompt)
    if tok > batas_prompt and LLM_ANALISIS_MODE != "penuh":
        prompt_kecil = _susun_prompt_penuh(system_prompt, teks_klaim, items, batas_chunk=min(TARIF_CHUNK_CHARS, 200), top_k=2)
        tok_kecil = _taksir_token(prompt_kecil)
        print(f"      [INFO] Prompt ~{tok:,} token > batas ~{batas_prompt:,}; dicoba versi ringkas (tarif 200 karakter, 2 kandidat): ~{tok_kecil:,} token.")
        if tok_kecil <= batas_prompt:
            prompt, tok = prompt_kecil, tok_kecil
    muat = tok <= batas_prompt
    print(f"      [INFO] Prompt analisis penuh ~{tok:,} token; batas ~{batas_prompt:,} (context {LLM_MAX_CONTEXT:,}, output {MAX_OUTPUT_ANALISIS:,}) "
          f"-> {'PENUH' if muat else 'tidak muat'}")

    if LLM_ANALISIS_MODE == "penuh" or (LLM_ANALISIS_MODE == "auto" and muat):
        hasil = _analisis_klaim_penuh(client, system_prompt, teks_klaim, items, prompt=prompt)
        return validasi_hasil(hasil, items, items_agregat)

    print(f"      [INFO] Memakai analisis berjenjang ({ANALISIS_BATCH_ITEM} item per panggilan).")
    hasil = _analisis_klaim_berjenjang(client, system_prompt, teks_klaim, items, gambar_halaman=gambar_halaman)
    hasil = validasi_hasil(hasil, items, items_agregat, hitung_ulang=True)
    _tulis_narasi(client, hasil, system_prompt, teks_klaim)
    return hasil


# ---------------------------------------------------------------------------
# Output: Markdown + Excel
# ---------------------------------------------------------------------------

# Karakter kontrol yang tidak valid di XML/Excel (sering muncul dari hasil OCR/vision)
_ILLEGAL_XLSX_CHARS = re.compile(r"[\000-\010\013\014\016-\037]")


def sel(value):
    """Bersihkan satu nilai sebelum ditulis ke sel Excel: buang karakter ilegal, batasi panjang."""
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    text = _ILLEGAL_XLSX_CHARS.sub("", text)
    return text[:32000]  # batas aman panjang sel Excel


def angka(value, default=0):
    """Paksa jadi angka dengan aman. Mengerti format Indonesia ('1.234.567,89') maupun US ('1,234,567.89')."""
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return value
    s = re.sub(r"[^\d.,\-]", "", str(value).strip())
    if not s or not re.search(r"\d", s):
        return default
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s):      # 1.234.567,89
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", s):    # 1,234,567.89
        s = s.replace(",", "")
    elif "," in s and "." not in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return default


PEMISAH_RIBUAN = ","  # ganti "." untuk gaya Indonesia: Rp 1.126.385.372


def rupiah(nilai) -> str:
    """Rupiah bulat tanpa desimal, mis. Rp 1,126,385,372. Nilai None/teks aneh dianggap 0."""
    n = int(round(angka(nilai, 0) or 0))
    teks = f"{abs(n):,}"
    if PEMISAH_RIBUAN != ",":
        teks = teks.replace(",", PEMISAH_RIBUAN)
    return f"Rp {'-' if n < 0 else ''}{teks}"


def _selisih_item(item: dict, harga, tarif):
    """Selisih yang sudah dihitung validasi_hasil kalau ada; kalau tidak, harga - tarif."""
    if tarif is None:
        return None
    s = item.get("selisih")
    return s if isinstance(s, (int, float)) else harga - tarif


_KUNCI_IDENTITAS = ("nama_pasien", "no_peserta", "rumah_sakit", "periode_rawat", "los_hari", "total_tagihan_rs", "program")


def _identitas_aman(hasil: dict) -> dict:
    """
    Baca blok identitas dengan toleran: kunci boleh huruf besar/spasi, dan kalau model lupa
    membungkus dengan "identitas", kunci di level atas dipakai. Nilai kosong ditampilkan "-".
    """
    sumber = hasil.get("identitas")
    if not isinstance(sumber, dict):
        sumber = {k: hasil.get(k) for k in _KUNCI_IDENTITAS}
    peta = {str(k).strip().lower().replace(" ", "_"): v for k, v in sumber.items()}
    out = {}
    for k in _KUNCI_IDENTITAS:
        v = peta.get(k)
        kosong = v is None or str(v).strip() == ""
        out[k] = (0 if k == "total_tagihan_rs" else "-") if kosong else v
    return out


def tulis_laporan_markdown(hasil: dict, out_path: Path) -> None:
    id_ = _identitas_aman(hasil)
    kp = hasil.get("keputusan_akhir", {})

    peringatan = hasil.get("peringatan") or []
    blok_peringatan = (["## ⚠️ PERINGATAN VALIDASI OTOMATIS\n"] + [f"- {p}" for p in peringatan] + ["\n---\n"]) if peringatan else []

    md = [
        "# ANALISIS DAN VERIFIKASI KLAIM KESEHATAN & JKK\n",
        f"**Nama Pasien:** {id_['nama_pasien']}",
        f"**No. Peserta:** {id_['no_peserta']}",
        f"**Rumah Sakit:** {id_['rumah_sakit']}",
        f"**Periode Rawat:** {id_['periode_rawat']}",
        f"**LOS:** {str(id_['los_hari']) + ' hari' if id_['los_hari'] != '-' else '-'}",
        f"**Total Tagihan RS:** {rupiah(angka(id_.get('total_tagihan_rs')))}",
        f"**Program:** {id_['program']}\n",
        "---\n",
        *blok_peringatan,
        hasil.get("narasi_lengkap", "(narasi tidak tersedia)"),
        "\n---\n",
        "## KEPUTUSAN AKHIR\n",
        f"**Status:** {kp.get('status', '')}\n",
        f"- Total Diajukan: {rupiah(angka(kp.get('total_diajukan')))}",
        f"- Koreksi Terverifikasi: {rupiah(angka(kp.get('koreksi_terverifikasi')))}",
        f"- Perlu Klarifikasi: {rupiah(angka(kp.get('perlu_klarifikasi')))}",
        f"- Rekomendasi Pembayaran: {rupiah(angka(kp.get('rekomendasi_pembayaran')))}\n",
        "**Alasan:**",
    ]
    for alasan in kp.get("alasan", []):
        md.append(f"- {alasan}")
    md.append(f"\n**Confidence Score:** {hasil.get('confidence_score', 0)}%")
    md.append(
        "\n\n> ⚠️ Ini adalah hasil decision-support otomatis. Keputusan akhir "
        "tetap wajib direview dan disahkan oleh verifikator manusia."
    )

    out_path.write_text("\n".join(md), encoding="utf-8")


def tulis_laporan_excel(hasil: dict, out_path: Path) -> None:
    wb = Workbook()

    # --- Sheet 1: Rincian Item ---
    ws1 = wb.active
    ws1.title = "Rincian Item"
    header = ["Kategori", "Item", "Tanggal", "Qty", "Harga Ditagih", "Tarif RS", "Selisih", "Keterangan", "Sumber Tarif"]
    ws1.append(header)
    for cell in ws1[1]:
        cell.font = Font(bold=True)

    total_ditagih = 0
    total_tarif = 0
    total_selisih = 0

    for item in hasil.get("tabel_item", []):
        harga = angka(item.get("harga_ditagih"))
        tarif_raw = item.get("tarif_rs")
        tarif = angka(tarif_raw) if tarif_raw is not None else None
        selisih = _selisih_item(item, harga, tarif)

        ws1.append([
            sel(item.get("kategori", "")),
            sel(item.get("item", "")),
            sel(item.get("tanggal", "")),
            sel(item.get("quantity", "")),
            harga,
            tarif if tarif is not None else "N/A",
            selisih if selisih is not None else "N/A",
            sel(item.get("keterangan", "")),
            sel(item.get("sumber_tarif", "")),
        ])
        total_ditagih += harga
        if tarif is not None:
            total_tarif += tarif
            total_selisih += selisih

    total_row = ["", "", "", "TOTAL", total_ditagih, total_tarif, total_selisih, "", ""]
    ws1.append(total_row)
    for cell in ws1[ws1.max_row]:
        cell.font = Font(bold=True)

    for col in ws1.columns:
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws1.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    # --- Sheet 2: Rekap Temuan ---
    ws2 = wb.create_sheet("Rekap Temuan")
    header2 = ["Jenis Temuan", "Jumlah Item", "Nilai (Rp)", "Severity", "Rekomendasi"]
    ws2.append(header2)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for t in hasil.get("rekap_temuan", []):
        ws2.append([
            sel(t.get("jenis_temuan", "")),
            sel(t.get("jumlah_item", "")),
            angka(t.get("nilai")),
            sel(t.get("severity", "")),
            sel(t.get("rekomendasi", "")),
        ])
    for col in ws2.columns:
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws2.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    if hasil.get("peringatan"):
        ws3 = wb.create_sheet("Peringatan")
        ws3.append(["Peringatan validasi otomatis"])
        ws3[1][0].font = Font(bold=True)
        for p in hasil["peringatan"]:
            ws3.append([sel(p)])
        ws3.column_dimensions["A"].width = 120

    wb.save(out_path)


def tulis_laporan_csv(hasil: dict, out_path: Path) -> None:
    """Cadangan selain .xlsx -- CSV jauh lebih jarang bermasalah dibuka di berbagai aplikasi."""
    with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["Kategori", "Item", "Tanggal", "Qty", "Harga Ditagih", "Tarif RS", "Selisih", "Keterangan", "Sumber Tarif"])

        total_ditagih = total_tarif = total_selisih = 0
        for item in hasil.get("tabel_item", []):
            harga = angka(item.get("harga_ditagih"))
            tarif_raw = item.get("tarif_rs")
            tarif = angka(tarif_raw) if tarif_raw is not None else None
            selisih = _selisih_item(item, harga, tarif)

            writer.writerow([
                sel(item.get("kategori", "")), sel(item.get("item", "")), sel(item.get("tanggal", "")),
                sel(item.get("quantity", "")), harga, tarif if tarif is not None else "N/A",
                selisih if selisih is not None else "N/A", sel(item.get("keterangan", "")), sel(item.get("sumber_tarif", "")),
            ])
            total_ditagih += harga
            if tarif is not None:
                total_tarif += tarif
                total_selisih += selisih

        writer.writerow(["", "", "", "TOTAL", total_ditagih, total_tarif, total_selisih, "", ""])


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--kode-rs", required=True, help="Kode rumah sakit (harus ada di tabel rumahsakit & sudah punya dokumen tarif di dokumen_chunk)")
    parser.add_argument("--claim-docs", required=True, nargs="+", help="Satu atau lebih file PDF dokumen klaim (invoice, resume medis, PLKK, dst)")
    parser.add_argument("--dsn", required=True, help="Connection string PostgreSQL")
    parser.add_argument("--llm-base-url", default=os.environ.get("LLM_BASE_URL"), help="URL tunnel vLLM, contoh: https://xxxx.trycloudflare.com/v1 (default: env LLM_BASE_URL)")
    parser.add_argument("--llm-api-key", default=os.environ.get("LLM_API_KEY", "EMPTY"), help="API key vLLM, isi EMPTY kalau server tanpa auth (default: env LLM_API_KEY)")
    parser.add_argument("--out-dir", default="./laporan_klaim", help="Folder output laporan")
    parser.add_argument("--resume-cache", default=None, help="Path ke file cache _items_*.json dari run sebelumnya, lanjut langsung ke tahap analisis (skip baca PDF + ekstraksi + vector search)")
    args = parser.parse_args()

    if not args.llm_base_url:
        sys.exit("URL vLLM tidak ditemukan. Isi LLM_BASE_URL di file .env, atau pakai --llm-base-url.")

    if not SYSTEM_PROMPT_PATH.exists():
        sys.exit(f"File system prompt tidak ditemukan: {SYSTEM_PROMPT_PATH}")

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    client = buat_llm_client(args.llm_base_url, args.llm_api_key)
    try:
        cek_server_llm(client)
    except RuntimeError as e:
        sys.exit(str(e))
    system_prompt = SYSTEM_PROMPT_PATH.read_text(encoding="utf-8")

    stem = f"laporan_{args.kode_rs}_{date.today().isoformat()}"
    cache_path = out_dir / f"_cache_items_{args.kode_rs}_{date.today().isoformat()}.json"

    if args.resume_cache:
        print(f"[1-3/5] Lanjut dari cache: {args.resume_cache} (skip baca PDF + ekstraksi + vector search)")
        with open(args.resume_cache, encoding="utf-8") as f:
            items = json.load(f)
        print(f"      -> {len(items)} item dimuat dari cache")
    else:
        print("[1/5] Membaca dokumen klaim")
        teks_klaim, gambar_halaman = baca_semua_dokumen_klaim(args.claim_docs)

        print(f"[2/5] Ekstraksi item tagihan (model: {MODEL_EKSTRAKSI})")
        items = ekstrak_item_tagihan(client, teks_klaim, gambar_halaman)
        print(f"      -> {len(items)} item tagihan ditemukan")

        print("[3/5] Mencari tarif pembanding di database (vector search)")
        embed_model = SentenceTransformer(EMBEDDING_MODEL_NAME)
        conn = psycopg2.connect(args.dsn)
        try:
            items = lengkapi_item_dengan_tarif(conn, embed_model, args.kode_rs, items)
        finally:
            conn.close()

        # Simpan cache SEBELUM masuk ke tahap analisis akhir yang paling sering gagal
        # kalau koneksi ke server LLM putus-nyambung -- supaya tidak perlu ulang dari awal.
        cache_path.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"      [cache tersimpan: {cache_path} -- kalau tahap analisis gagal, lanjut pakai --resume-cache {cache_path}]")

        # Simpan juga teks klaim mentahnya, dibutuhkan tahap analisis
        (out_dir / f"_cache_teks_{args.kode_rs}_{date.today().isoformat()}.txt").write_text(teks_klaim, encoding="utf-8")

    if args.resume_cache:
        teks_path = Path(args.resume_cache).parent / f"_cache_teks_{args.kode_rs}_{date.today().isoformat()}.txt"
        teks_klaim = teks_path.read_text(encoding="utf-8") if teks_path.exists() else ""
        gambar_halaman = []

    print(f"[4/5] Analisis akhir (model: {MODEL_ANALISIS})")
    hasil = analisis_klaim(client, system_prompt, teks_klaim, items, gambar_halaman=gambar_halaman)

    print("[5/5] Menyimpan laporan")
    md_path = out_dir / f"{stem}.md"
    xlsx_path = out_dir / f"{stem}.xlsx"
    csv_path = out_dir / f"{stem}.csv"
    tulis_laporan_markdown(hasil, md_path)
    tulis_laporan_excel(hasil, xlsx_path)
    tulis_laporan_csv(hasil, csv_path)

    kp = hasil.get("keputusan_akhir", {})
    jumlah_tabel = len(hasil.get("tabel_item", []))
    n_item_normal = len([i for i in items if not i.get("_agregat")])
    print(f"\n      -> {md_path}")
    print(f"      -> {xlsx_path}")
    print(f"      -> {csv_path}  (cadangan, buka ini kalau xlsx bermasalah)")
    print(f"\nJumlah item di tabel akhir: {jumlah_tabel} (dari {n_item_normal} item non-subtotal yang diekstrak)")
    if jumlah_tabel != n_item_normal:
        print("⚠️  JUMLAH TIDAK COCOK -- kemungkinan ada item yang hilang/diringkas saat analisis. Cek manual.")
    print(f"Status: {kp.get('status', '(tidak ada)')}")
    print(f"Confidence: {hasil.get('confidence_score', '?')}%")
    print("\n⚠️  Hasil ini decision-support, tetap wajib direview verifikator manusia.")


if __name__ == "__main__":
    main()